from pathlib import Path
import json
import sys

DESKTOP_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(DESKTOP_DIR))

import vtt_core as vc


def seg(i, start, end, text, speaker=None, log=-0.2):
    return {
        "id": i, "start": start, "end": end, "text": text,
        "speaker": speaker, "avg_logprob": log,
        "no_speech_prob": 0.05, "compression_ratio": 1.1, "words": [],
    }


def test_agrupar_segmentos_reduce_fragmentacion():
    segmentos = [
        seg(0, 0, 3, "Hola, este es el primer segmento."),
        seg(1, 3.1, 6, "Sigue la misma idea."),
        seg(2, 6.1, 9, "Y termina aquí."),
    ]
    bloques = vc.agrupar_segmentos(segmentos, min_caracteres_oracion=500)
    assert len(bloques) == 1
    assert bloques[0]["start"] == 0
    assert bloques[0]["end"] == 9
    assert "primer segmento" in bloques[0]["text"]


def test_agrupar_corta_por_hablante():
    segmentos = [
        seg(0, 0, 3, "Pregunta", "Persona 1"),
        seg(1, 3, 6, "Respuesta", "Persona 2"),
    ]
    bloques = vc.agrupar_segmentos(segmentos)
    assert [b["speaker"] for b in bloques] == ["Persona 1", "Persona 2"]


def test_agrupar_corta_por_pausa():
    segmentos = [seg(0, 0, 2, "Uno"), seg(1, 4, 6, "Dos")]
    assert len(vc.agrupar_segmentos(segmentos, pausa_corte=1.4)) == 2


def test_asignar_hablante_por_mayor_solapamiento():
    segmentos = [seg(0, 0, 5, "texto"), seg(1, 5, 10, "otro")]
    turnos = [
        {"start": 0, "end": 4.8, "speaker": 7},
        {"start": 5.1, "end": 10, "speaker": 9},
    ]
    salida, speakers = vc.asignar_hablantes(segmentos, turnos)
    assert salida[0]["speaker"] == "Persona 1"
    assert salida[1]["speaker"] == "Persona 2"
    assert len(speakers) == 2


def test_marca_baja_confianza_sin_borrar_texto():
    s = seg(0, 0, 2, "texto dudoso", log=-1.5)
    bloques = vc.agrupar_segmentos([s])
    assert bloques[0]["review_required"] is True
    assert bloques[0]["text"] == "texto dudoso"


def test_metricas_rtf_y_velocidad():
    m = vc.calcular_metricas(600, 120)
    assert m["rtf"] == 0.2
    assert m["speed_x"] == 5.0


def test_json_schema_2_con_bloques_y_metricas():
    segmentos = [seg(0, 0, 2, "hola")]
    bloques = vc.agrupar_segmentos(segmentos)
    doc = vc.documento_json("a.wav", "small", "es", "hola", segmentos, bloques, [],
                            vc.calcular_metricas(2, 1), "Equilibrado", "SITFA; Concepción")
    texto = vc.json_texto(doc)
    data = json.loads(texto)
    assert data["schema_version"] == 2
    assert data["reading_blocks"]
    assert data["hotwords"] == ["SITFA", "Concepción"]


def test_glosario_elimina_duplicados_preservando_orden():
    assert vc.glosario_a_hotwords("SITFA; RIT; SITFA") == "SITFA, RIT"


def test_srt_conserva_segmentos_tecnicos_y_hablante():
    s = seg(0, 0, 2.2, "Hola", "Persona 1")
    out = vc.srt_segmentos([s])
    assert "00:00:00,000 --> 00:00:02,200" in out
    assert "Persona 1: Hola" in out


def test_perfiles_tienen_modo_rapido_equilibrado_preciso():
    assert set(vc.PERFILES) == {"Rapido", "Equilibrado", "Preciso"}
    assert vc.PERFILES["Rapido"]["beam_size"] == 1
    assert vc.PERFILES["Equilibrado"]["batched"] is True
    assert vc.PERFILES["Preciso"]["batched"] is False


def test_docx_se_genera_y_reabre(tmp_path):
    from docx import Document
    bloques = vc.agrupar_segmentos([seg(0, 0, 3, "Hola mundo.", "Persona 1")])
    ruta = tmp_path / "salida.docx"
    vc.escribir_docx(ruta, "audio.wav", "small", "es", bloques, vc.calcular_metricas(3, 1))
    assert ruta.is_file() and ruta.stat().st_size > 1000
    doc = Document(ruta)
    assert doc.paragraphs[0].text == "TRANSCRIPCIÓN"
    assert any("Persona 1" in p.text for p in doc.paragraphs)
