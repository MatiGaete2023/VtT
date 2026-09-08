#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reportes V4 con diagnóstico de diarización."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Mapping, Sequence

import vtt_core as core
import vtt_reporting as v3


def documento_json_detallado(
    archivo: str,
    modelo: str,
    idioma: str,
    texto: str,
    segmentos: Sequence[Mapping[str, Any]],
    bloques: Sequence[Mapping[str, Any]],
    speakers: Sequence[Mapping[str, Any]],
    metricas: Mapping[str, Any],
    perfil: str,
    glosario: str,
) -> Dict[str, Any]:
    doc = v3.documento_json_detallado(
        archivo, modelo, idioma, texto, segmentos, bloques, speakers,
        metricas, perfil, glosario,
    )
    doc["schema_version"] = max(4, int(doc.get("schema_version", 0) or 0))
    cfg = doc.setdefault("configuration", {})
    cfg["diarization_profile"] = metricas.get("diarization_profile")
    cfg["window_shift_ratio"] = metricas.get("window_shift_ratio")

    diar = doc.setdefault("diarization", {})
    diar.update({
        "profile": metricas.get("diarization_profile"),
        "window_shift_ratio": metricas.get("window_shift_ratio"),
        "passes": metricas.get("auto_passes", 1),
        "selection_reason": metricas.get("auto_selection_reason"),
        "retry_reason": metricas.get("auto_retry_reason"),
        "stability_delta_speakers": metricas.get("auto_stability_delta_speakers"),
        "selected_analysis": metricas.get("auto_selected_analysis"),
        "candidates": metricas.get("auto_candidates") or [],
    })
    return doc


def _segundos(metricas: Mapping[str, Any], clave: str) -> float:
    try:
        return max(0.0, float(metricas.get(clave, 0.0) or 0.0))
    except (TypeError, ValueError):
        return 0.0


def escribir_docx_detallado(
    ruta: Path,
    archivo: str,
    modelo: str,
    idioma: str,
    bloques: Sequence[Mapping[str, Any]],
    metricas: Mapping[str, Any],
    perfil: str,
) -> None:
    from docx import Document
    from docx.shared import Pt

    m = v3.cerrar_metricas(metricas)
    doc = Document()
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("TRANSCRIPCIÓN", level=0)
    tabla = doc.add_table(rows=0, cols=2)

    backend = "—"
    if m.get("device"):
        backend = str(m.get("device"))
        if m.get("compute_type"):
            backend += f" / {m.get('compute_type')}"

    if not m.get("diarization_enabled"):
        hablantes = "No"
    elif m.get("speaker_mode") == "Auto":
        hablantes = f"Auto · detectados: {int(m.get('speaker_detected', 0) or 0)}"
    else:
        hablantes = (
            f"Solicitados: {m.get('speaker_requested')} · "
            f"detectados: {int(m.get('speaker_detected', 0) or 0)}"
        )

    datos = [
        ("Archivo", Path(archivo).name),
        ("Modelo", modelo),
        ("Idioma", idioma or "auto"),
        ("Perfil ASR", perfil),
        ("Backend", backend),
        ("Batch", m.get("batch_size") if m.get("batched") else "secuencial"),
        ("Beam", m.get("beam_size", "—")),
        ("Hablantes", hablantes),
        ("Perfil diarización", m.get("diarization_profile") or "—"),
        (
            "Window shift",
            f"{float(m['window_shift_ratio']):.2f}"
            if m.get("window_shift_ratio") is not None else "—",
        ),
        ("Hilos diarización", m.get("diarization_threads") or "—"),
        ("Pasadas Auto", m.get("auto_passes") or 1),
        ("Selección Auto", m.get("auto_selection_reason") or "—"),
        ("Duración", core.ts_simple(_segundos(m, "audio_seconds"))),
        ("Carga de modelo", core.ts_simple(_segundos(m, "model_load_seconds"))),
        ("Transcripción ASR", core.ts_simple(_segundos(m, "asr_seconds"))),
        ("Identificación hablantes", core.ts_simple(_segundos(m, "diarization_seconds"))),
        ("Exportación", core.ts_simple(_segundos(m, "export_seconds"))),
        ("Otros", core.ts_simple(_segundos(m, "overhead_seconds"))),
        ("Procesamiento", core.ts_simple(_segundos(m, "processing_seconds"))),
        ("Velocidad", f"{float(m.get('speed_x', 0.0) or 0.0):.2f}× tiempo real"),
    ]
    if m.get("auto_threshold") is not None:
        datos.append(("Auto hablantes", f"umbral elegido {float(m['auto_threshold']):.2f}"))
    if m.get("auto_retry_reason"):
        datos.append(("Motivo segunda pasada", m.get("auto_retry_reason")))
    reduccion = m.get("speech_region_reduction") or {}
    if isinstance(reduccion, Mapping) and reduccion.get("enabled"):
        datos.append((
            "Regiones de voz",
            f"{float(reduccion.get('coverage', 0.0)) * 100:.0f}% del audio procesado",
        ))

    for k, v in datos:
        c = tabla.add_row().cells
        c[0].text = str(k)
        c[1].text = str(v)

    doc.add_paragraph()
    for b in bloques:
        p = doc.add_paragraph()
        run = p.add_run(f"[{core.ts_simple(float(b.get('start', 0.0)))}]")
        run.bold = True
        if b.get("speaker"):
            r = p.add_run(f"  {b['speaker']}")
            r.bold = True
        if b.get("review_required"):
            r = p.add_run("  [REVISAR]")
            r.bold = True
        doc.add_paragraph(core.limpiar_texto(str(b.get("text", ""))))
    doc.save(str(ruta))
