#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reportes V5.2: conteos separados, identidad interpretable y rendimiento."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Mapping, Sequence

from docx import Document
from docx.shared import Pt

import vtt_core as core
import vtt_performance as perf
import vtt_reporting as base_reporting
import vtt_reporting_v51 as v51
import vtt_validation_v52 as validation52


def documento_json_detallado(
    archivo: str, modelo: str, idioma: str, texto: str,
    segmentos: Sequence[Mapping[str, Any]], bloques: Sequence[Mapping[str, Any]],
    speakers: Sequence[Mapping[str, Any]], metricas: Mapping[str, Any],
    perfil: str, glosario: str,
) -> Dict[str, Any]:
    # Nunca reutilizar un estado ``performance`` precalculado: exportación y
    # reportes pueden haber cambiado los tiempos desde _metricas_base().
    m = base_reporting.cerrar_metricas(metricas)
    m["performance"] = perf.performance_status(m)
    doc = v51.documento_json_detallado(
        archivo, modelo, idioma, texto, segmentos, bloques,
        speakers, m, perfil, glosario,
    )
    doc["schema_version"] = max(8, int(doc.get("schema_version", 0) or 0))
    diar = doc.setdefault("diarization", {})
    diar["speaker_counts"] = m.get("speaker_counts") or {}
    diar["auto_precheck"] = m.get("auto_precheck") or {}
    diar["identity_aware_selection"] = m.get("identity_aware_selection") or {}
    diar["retry_requested"] = bool(m.get("auto_retry_requested", False))
    diar["retry_avoided"] = bool(m.get("auto_retry_avoided", False))
    diar["retry_skipped_budget"] = bool(m.get("auto_retry_skipped_budget", False))
    diar["time_budget_seconds"] = m.get("diarization_time_budget_seconds")
    doc["performance"] = perf.performance_status(m)
    doc["global_profile"] = m.get("global_profile")
    return doc


def _fmt_s(value: Any) -> str:
    try:
        return f"{float(value or 0.0):.2f} s"
    except (TypeError, ValueError):
        return "—"


def _fmt_clock(value: Any) -> str:
    try:
        return core.ts_simple(max(0.0, float(value or 0.0)))
    except (TypeError, ValueError):
        return "—"


def _fmt_identity(sp: Mapping[str, Any]) -> str:
    confidence = str(sp.get("confidence", "sin_datos"))
    turns = int(sp.get("turns", 0) or 0)
    embedded = int(sp.get("embedded_turns", 0) or 0)
    seconds = float(sp.get("total_seconds", 0.0) or 0.0)
    gap = float(sp.get("max_gap_seconds", 0.0) or 0.0)
    base = f"{confidence} · {turns} turno(s) · {seconds:.1f} s · {embedded} muestra(s) acústica(s)"
    if embedded == 0:
        return base + " · similitud no evaluable"
    if embedded == 1:
        return base + " · muestra única; similitud no evaluable"
    if embedded == 2:
        pair = sp.get("pair_similarity"); rival = sp.get("rival_similarity_max")
        if pair is not None:
            base += f" · similitud entre apariciones {float(pair):.2f}"
        if rival is not None:
            base += f" · rival máx {float(rival):.2f}"
        if gap > 0:
            base += f" · separación máx {gap:.1f} s"
        return base
    med = sp.get("prototype_similarity_median"); minv = sp.get("prototype_similarity_min")
    rival = sp.get("rival_similarity_max")
    if med is not None:
        base += f" · prototipo mediana {float(med):.2f}"
    if minv is not None:
        base += f" · mínimo {float(minv):.2f}"
    if rival is not None:
        base += f" · rival máx {float(rival):.2f}"
    return base


def escribir_docx_detallado(
    ruta: Path, archivo: str, modelo: str, idioma: str,
    bloques: Sequence[Mapping[str, Any]], metricas: Mapping[str, Any], perfil: str,
) -> None:
    """Construye el DOCX V5.2 directamente y lo guarda una sola vez."""
    m = base_reporting.cerrar_metricas(metricas)
    perf_status = perf.performance_status(m)
    counts = m.get("speaker_counts") or {}
    valid = m.get("speaker_count_validation") or {}
    acoustic = int(counts.get("raw_acoustic_clusters", m.get("speaker_detected", 0)) or 0)
    identity_n = int(counts.get("identity_clusters_after_refinement", acoustic) or 0)
    text_n = int(counts.get("text_assigned_speakers", m.get("speaker_text_assigned", 0)) or 0)
    unassigned = set(int(x) for x in (counts.get("unassigned_raw_ids") or []))

    doc = Document()
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("TRANSCRIPCIÓN", level=0)
    table = doc.add_table(rows=0, cols=2)

    backend = "—"
    if m.get("device"):
        backend = str(m.get("device"))
        if m.get("compute_type"):
            backend += f" / {m.get('compute_type')}"

    if not m.get("diarization_enabled"):
        hablantes = "No"
    elif m.get("speaker_mode") == "Auto":
        hablantes = (
            f"Auto · estimación acústica: {identity_n} · con texto: {text_n} · "
            f"{validation52.etiqueta(valid)}"
        )
    else:
        hablantes = (
            f"Solicitados: {m.get('speaker_requested')} · acústicos: {identity_n} · "
            f"con texto: {text_n}"
        )

    rows = [
        ("Archivo", Path(archivo).name),
        ("Modelo", modelo),
        ("Idioma", idioma or "auto"),
        ("Perfil ASR", perfil),
        ("Backend", backend),
        ("Batch", m.get("batch_size") if m.get("batched") else "secuencial"),
        ("Beam", m.get("beam_size", "—")),
        ("Hablantes", hablantes),
        ("Perfil diarización", m.get("diarization_profile") or "—"),
        ("Window shift", f"{float(m['window_shift_ratio']):.2f}" if m.get("window_shift_ratio") is not None else "—"),
        ("Hilos diarización", m.get("diarization_threads") or "—"),
        ("Pasadas Auto", m.get("auto_passes") or 1),
        ("Selección Auto", m.get("auto_selection_reason") or "—"),
        ("Duración", _fmt_clock(m.get("audio_seconds"))),
        ("Carga de modelo", _fmt_clock(m.get("model_load_seconds"))),
        ("Transcripción ASR", _fmt_clock(m.get("asr_seconds"))),
        ("Identificación hablantes", _fmt_clock(m.get("diarization_seconds"))),
        ("Exportación", _fmt_clock(m.get("export_seconds"))),
        ("Otros", _fmt_clock(m.get("overhead_seconds"))),
        ("Procesamiento", _fmt_clock(m.get("processing_seconds"))),
        ("Espera extremo a extremo", _fmt_clock(m.get("end_to_end_seconds"))),
        ("Velocidad", f"{float(m.get('speed_x', 0.0) or 0.0):.2f}× tiempo real"),
    ]
    if m.get("auto_threshold") is not None:
        rows.append(("Auto hablantes", f"umbral elegido {float(m['auto_threshold']):.2f}"))
    if m.get("auto_retry_reason"):
        rows.append(("Motivo segunda pasada", str(m.get("auto_retry_reason"))))

    align = m.get("speaker_alignment") or {}
    rows.extend([
        ("Alineación hablantes", f"{align.get('mode', '—')} · {int(align.get('split_source_segments', 0) or 0)} segmento(s) divididos · {int(align.get('speaker_switches_inside_segments', 0) or 0)} cambio(s) interno(s)"),
        ("Timestamps palabra internos", "Sí" if m.get("word_timestamps_forced_for_diarization") else "No"),
        ("Worker diarización", f"persistente · job {m.get('diarization_worker_job_index') or '—'} · motor {'reutilizado' if m.get('diarization_engine_reused') else 'inicializado'}"),
        ("Preparación modelos diar.", _fmt_s(m.get("diarization_model_prepare_seconds"))),
        ("Inicialización motor diar.", _fmt_s(m.get("diarization_engine_init_seconds"))),
        ("Decodificación diar.", _fmt_s(m.get("diarization_decode_seconds"))),
        ("Proceso sherpa (wall)", _fmt_s(m.get("diarization_process_wall_seconds"))),
    ])
    internal = m.get("sherpa_internal") or {}
    if m.get("sherpa_internal_timing_available"):
        rows.extend([
            ("Sherpa segmentación", _fmt_s(internal.get("segmentation_seconds"))),
            ("Sherpa embeddings", _fmt_s(internal.get("embedding_seconds"))),
            ("Sherpa clustering", _fmt_s(internal.get("clustering_seconds"))),
            ("Sherpa total interno", _fmt_s(internal.get("sherpa_total_seconds"))),
        ])
    else:
        rows.append(("Timers internos sherpa", "no capturados por el backend/SO; se conserva medición wall"))

    if valid:
        rows.append((
            "Validación conteo hablantes",
            f"{valid.get('status', '—')} · acústicos {identity_n} · con texto {text_n} · "
            f"ground truth {'sí' if valid.get('ground_truth_available') else 'no'} · "
            f"identidad {valid.get('identity_confidence', 'sin_datos')}",
        ))

    for p in m.get("sherpa_pass_timings") or []:
        rows.append((f"Sherpa pasada {int(p.get('pass', 0) or 0)} (wall)", _fmt_s(p.get("process_wall_seconds"))))

    iv = m.get("identity_verification") or {}
    consistency = m.get("identity_consistency") or {}
    refinement = m.get("identity_refinement") or {}
    scan = m.get("identity_local_scan") or {}
    if iv.get("enabled"):
        rows.extend([
            ("Control identidad V5.2", "Sí · control acústico conservador y selección identity-aware"),
            ("Identidad total (wall)", _fmt_s(iv.get("total_wall_seconds", m.get("identity_wall_seconds", 0.0)))),
            ("Embeddings identidad", f"{int(m.get('identity_embedding_calls', 0) or 0)} cálculo(s) · {_fmt_s(m.get('identity_embedding_compute_seconds'))}"),
            ("Refinamiento identidad", f"{int(refinement.get('reassigned_turns', 0) or 0)} turno(s) reasignado(s) · {int(refinement.get('new_identities', 0) or 0)} identidad(es) nueva(s)"),
            ("Escaneo local", f"{int(scan.get('scanned_turns', 0) or 0)} turno(s) largo(s) · {int(scan.get('applied_changes', 0) or 0)} cambio(s) aplicado(s)"),
            ("Consistencia identidad", str(consistency.get("overall_confidence", "sin_datos"))),
        ])
        if "final_stage_wall_seconds" in iv:
            rows.append(("Identidad etapa final (wall)", _fmt_s(iv.get("final_stage_wall_seconds"))))
        if "light_identity_wall_seconds" in iv:
            rows.append(("Identidad ligera (wall)", _fmt_s(iv.get("light_identity_wall_seconds"))))
        for sp in consistency.get("speakers", []) or []:
            raw = int(sp.get("speaker", -1))
            display = sp.get("display_name")
            if display:
                name = f"Identidad {display}"
            elif raw in unassigned:
                name = f"Cluster acústico raw {raw} (sin texto)"
            else:
                name = f"Cluster acústico raw {raw}"
            rows.append((name, _fmt_identity(sp)))

    if str(m.get("device", "")) == "cpu" and m.get("diarization_profile") == "Precisa":
        rows.append(("Rendimiento diarización", "Precisa prioriza resolución temporal y puede ser muy lenta en CPU; Equilibrada es el perfil recomendado."))

    rows.extend([
        ("Clusters sherpa seleccionados", str(acoustic)),
        ("Clusters tras control identidad", str(identity_n)),
        ("Hablantes con texto", str(text_n)),
        ("Clusters acústicos sin texto", ", ".join(str(x) for x in sorted(unassigned)) if unassigned else "0"),
        ("Segunda pasada solicitada", "Sí" if m.get("auto_retry_requested") else "No"),
        ("Segunda pasada evitada", "Sí" if m.get("auto_retry_avoided") else "No"),
        ("Segunda pasada omitida por presupuesto", "Sí" if m.get("auto_retry_skipped_budget") else "No"),
    ])
    if m.get("diarization_time_budget_seconds") is not None:
        rows.append(("Presupuesto diarización Auto", _fmt_s(m.get("diarization_time_budget_seconds"))))

    pre = m.get("auto_precheck") or {}
    if pre.get("enabled"):
        rows.append((
            "Precheck rendimiento Auto",
            f"{float(pre.get('wall_seconds', 0.0) or 0.0):.2f} s · "
            f"{int(pre.get('speakers_before', 0) or 0)}→{int(pre.get('speakers_after', 0) or 0)} · "
            f"{'aceptado' if pre.get('accepted') else 'no concluyente'}",
        ))
    sel = m.get("identity_aware_selection") or {}
    if sel.get("enabled"):
        rows.append((
            "Selección Auto identity-aware",
            f"{sel.get('selected_candidate', '—')} · score 1 {float(sel.get('first_score', 0.0)):.2f} · "
            f"score 2 {float(sel.get('second_score', 0.0)):.2f}",
        ))

    if perf_status.get("within_realtime") is not None:
        rows.append((
            "Presupuesto de rendimiento",
            ("cumplido" if perf_status.get("within_realtime") else "sobre tiempo real")
            + f" · proceso/audio {float(perf_status.get('processing_to_audio', 0.0)):.2f}×",
        ))
    if m.get("global_profile"):
        rows.append(("Modo global", str(m.get("global_profile"))))

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)

    doc.add_paragraph()
    for b in bloques:
        p = doc.add_paragraph()
        run = p.add_run(f"[{core.ts_simple(float(b.get('start', 0.0)))}]")
        run.bold = True
        if b.get("speaker"):
            r = p.add_run(f"  {b['speaker']}"); r.bold = True
        if b.get("review_required"):
            r = p.add_run("  [REVISAR]"); r.bold = True
        doc.add_paragraph(core.limpiar_texto(str(b.get("text", ""))))

    doc.save(str(ruta))
