#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Exportación detallada para VtT.

Añade a JSON/DOCX metadatos suficientes para diagnosticar rendimiento y
configuración sin alterar el texto transcrito.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Mapping, Sequence

import vtt_core as core


def _segundos(metricas: Mapping[str, Any], clave: str) -> float:
    try:
        return max(0.0, float(metricas.get(clave, 0.0) or 0.0))
    except (TypeError, ValueError):
        return 0.0


def cerrar_metricas(metricas: Mapping[str, Any]) -> Dict[str, Any]:
    """Normaliza tiempos de etapas y recalcula total, RTF y velocidad."""
    m = dict(metricas)
    audio = _segundos(m, "audio_seconds")
    asr = _segundos(m, "asr_seconds")
    diar = _segundos(m, "diarization_seconds")
    export = _segundos(m, "export_seconds")
    overhead = _segundos(m, "overhead_seconds")
    model_load = _segundos(m, "model_load_seconds")
    total = asr + diar + export + overhead
    m["processing_seconds"] = total
    m["rtf"] = total / audio if audio > 0 else None
    m["speed_x"] = audio / total if total > 0 else None
    m["asr_seconds"] = asr
    m["diarization_seconds"] = diar
    m["export_seconds"] = export
    m["overhead_seconds"] = overhead
    m["model_load_seconds"] = model_load
    return m


def configuracion(metricas: Mapping[str, Any], perfil: str) -> Dict[str, Any]:
    return {
        "profile": perfil,
        "device": metricas.get("device"),
        "compute_type": metricas.get("compute_type"),
        "batched": bool(metricas.get("batched", False)),
        "batch_size": metricas.get("batch_size"),
        "beam_size": metricas.get("beam_size"),
        "diarization_enabled": bool(metricas.get("diarization_enabled", False)),
        "speaker_mode": metricas.get("speaker_mode", "No"),
        "speaker_requested": metricas.get("speaker_requested"),
        "speaker_detected": metricas.get("speaker_detected", 0),
        "auto_threshold": metricas.get("auto_threshold"),
        "speech_region_reduction": metricas.get("speech_region_reduction"),
        "diarization_threads": metricas.get("diarization_threads"),
        "auto_retry": metricas.get("auto_retry"),
    }


def documento_json_detallado(
    archivo: str,
    modelo: str,
    idioma: str,
    texto: str,
    segmentos: Sequence[Mapping[str, Any]],
    bloques: Sequence[Mapping[str, Any]],
    speakers: Sequence[Mapping[str, Any]],
    metricas: Mapping[str, Any],
    perfil: str,
    glosario: str,
) -> Dict[str, Any]:
    m = cerrar_metricas(metricas)
    doc = core.documento_json(
        archivo, modelo, idioma, texto, segmentos, bloques, speakers,
        m, perfil, glosario,
    )
    doc["schema_version"] = max(3, int(doc.get("schema_version", 0) or 0))
    doc["configuration"] = configuracion(m, perfil)
    doc["timing"] = {
        "model_load_seconds": m.get("model_load_seconds", 0.0),
        "asr_seconds": m.get("asr_seconds", 0.0),
        "diarization_seconds": m.get("diarization_seconds", 0.0),
        "export_seconds": m.get("export_seconds", 0.0),
        "overhead_seconds": m.get("overhead_seconds", 0.0),
        "processing_seconds": m.get("processing_seconds", 0.0),
        "rtf": m.get("rtf"),
        "speed_x": m.get("speed_x"),
    }
    doc["diarization"] = {
        "enabled": bool(m.get("diarization_enabled", False)),
        "mode": m.get("speaker_mode", "No"),
        "requested_speakers": m.get("speaker_requested"),
        "detected_speakers": m.get("speaker_detected", 0),
        "selected_threshold": m.get("auto_threshold"),
        "pilot": m.get("auto_pilot"),
        "speech_region_reduction": m.get("speech_region_reduction"),
        "diarization_threads": m.get("diarization_threads"),
        "auto_retry": m.get("auto_retry"),
        "initial_speakers": m.get("auto_initial_speakers"),
    }
    return doc


def escribir_docx_detallado(
    ruta: Path,
    archivo: str,
    modelo: str,
    idioma: str,
    bloques: Sequence[Mapping[str, Any]],
    metricas: Mapping[str, Any],
    perfil: str,
) -> None:
    """DOCX literal con configuración y tiempos por etapa."""
    from docx import Document
    from docx.shared import Pt

    m = cerrar_metricas(metricas)
    doc = Document()
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("TRANSCRIPCIÓN", level=0)
    tabla = doc.add_table(rows=0, cols=2)

    backend = "—"
    if m.get("device"):
        backend = str(m.get("device"))
        if m.get("compute_type"):
            backend += f" / {m.get('compute_type')}"

    if not m.get("diarization_enabled"):
        hablantes = "No"
    elif m.get("speaker_mode") == "Auto":
        hablantes = f"Auto · detectados: {int(m.get('speaker_detected', 0) or 0)}"
    else:
        hablantes = (
            f"Solicitados: {m.get('speaker_requested')} · "
            f"detectados: {int(m.get('speaker_detected', 0) or 0)}"
        )

    datos = [
        ("Archivo", Path(archivo).name),
        ("Modelo", modelo),
        ("Idioma", idioma or "auto"),
        ("Perfil", perfil),
        ("Backend", backend),
        ("Batch", m.get("batch_size") if m.get("batched") else "secuencial"),
        ("Beam", m.get("beam_size", "—")),
        ("Hablantes", hablantes),
        ("Hilos diarización", m.get("diarization_threads") or "—"),
        ("Duración", core.ts_simple(_segundos(m, "audio_seconds"))),
        ("Carga de modelo", core.ts_simple(_segundos(m, "model_load_seconds"))),
        ("Transcripción ASR", core.ts_simple(_segundos(m, "asr_seconds"))),
        ("Identificación hablantes", core.ts_simple(_segundos(m, "diarization_seconds"))),
        ("Exportación", core.ts_simple(_segundos(m, "export_seconds"))),
        ("Otros", core.ts_simple(_segundos(m, "overhead_seconds"))),
        ("Procesamiento", core.ts_simple(_segundos(m, "processing_seconds"))),
        ("Velocidad", f"{float(m.get('speed_x', 0.0) or 0.0):.2f}× tiempo real"),
    ]
    if m.get("auto_threshold") is not None:
        datos.append(("Auto hablantes", f"umbral elegido {float(m['auto_threshold']):.2f}"))
    reduccion = m.get("speech_region_reduction") or {}
    if isinstance(reduccion, Mapping) and reduccion.get("enabled"):
        datos.append((
            "Regiones de voz",
            f"{float(reduccion.get('coverage', 0.0)) * 100:.0f}% del audio procesado",
        ))

    for k, v in datos:
        c = tabla.add_row().cells
        c[0].text = str(k)
        c[1].text = str(v)

    doc.add_paragraph()
    for b in bloques:
        p = doc.add_paragraph()
        run = p.add_run(f"[{core.ts_simple(float(b.get('start', 0.0)))}]")
        run.bold = True
        if b.get("speaker"):
            r = p.add_run(f"  {b['speaker']}")
            r.bold = True
        if b.get("review_required"):
            r = p.add_run("  [REVISAR]")
            r.bold = True
        doc.add_paragraph(core.limpiar_texto(str(b.get("text", ""))))
    doc.save(str(ruta))
