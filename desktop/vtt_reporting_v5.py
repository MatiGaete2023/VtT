#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reportes V5: alineación por palabra, worker persistente y tiempos sherpa."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Mapping, Sequence

from docx import Document

import vtt_reporting_v4 as v4


def documento_json_detallado(
    archivo: str,
    modelo: str,
    idioma: str,
    texto: str,
    segmentos: Sequence[Mapping[str, Any]],
    bloques: Sequence[Mapping[str, Any]],
    speakers: Sequence[Mapping[str, Any]],
    metricas: Mapping[str, Any],
    perfil: str,
    glosario: str,
) -> Dict[str, Any]:
    doc = v4.documento_json_detallado(
        archivo, modelo, idioma, texto, segmentos, bloques, speakers,
        metricas, perfil, glosario,
    )
    doc["schema_version"] = max(5, int(doc.get("schema_version", 0) or 0))
    cfg = doc.setdefault("configuration", {})
    cfg["word_timestamps_forced_for_diarization"] = bool(
        metricas.get("word_timestamps_forced_for_diarization", False)
    )

    diar = doc.setdefault("diarization", {})
    diar["persistent_worker"] = {
        "job_index": metricas.get("diarization_worker_job_index"),
        "models_reused": metricas.get("diarization_models_reused"),
        "engine_reused": metricas.get("diarization_engine_reused"),
    }
    diar["timing"] = {
        "model_prepare_seconds": metricas.get("diarization_model_prepare_seconds", 0.0),
        "engine_init_seconds": metricas.get("diarization_engine_init_seconds", 0.0),
        "clustering_config_seconds": metricas.get("diarization_clustering_config_seconds", 0.0),
        "decode_seconds": metricas.get("diarization_decode_seconds", 0.0),
        "reduction_seconds": metricas.get("diarization_reduction_seconds", 0.0),
        "process_wall_seconds": metricas.get("diarization_process_wall_seconds", 0.0),
        "internal_available": bool(metricas.get("sherpa_internal_timing_available", False)),
        "sherpa_internal": metricas.get("sherpa_internal") or {},
        "passes": metricas.get("sherpa_pass_timings") or [],
    }
    diar["speaker_count_validation"] = metricas.get("speaker_count_validation") or {}
    doc["alignment"] = metricas.get("speaker_alignment") or {}
    return doc


def _fmt_s(v: Any) -> str:
    try:
        return f"{float(v or 0.0):.2f} s"
    except (TypeError, ValueError):
        return "—"


def escribir_docx_detallado(
    ruta: Path,
    archivo: str,
    modelo: str,
    idioma: str,
    bloques: Sequence[Mapping[str, Any]],
    metricas: Mapping[str, Any],
    perfil: str,
) -> None:
    v4.escribir_docx_detallado(ruta, archivo, modelo, idioma, bloques, metricas, perfil)
    doc = Document(str(ruta))
    if not doc.tables:
        doc.save(str(ruta))
        return
    tabla = doc.tables[0]

    align = metricas.get("speaker_alignment") or {}
    valid = metricas.get("speaker_count_validation") or {}
    internal = metricas.get("sherpa_internal") or {}
    rows = [
        ("Alineación hablantes", f"{align.get('mode', '—')} · {int(align.get('split_source_segments', 0) or 0)} segmento(s) divididos · {int(align.get('speaker_switches_inside_segments', 0) or 0)} cambio(s) interno(s)"),
        ("Timestamps palabra internos", "Sí" if metricas.get("word_timestamps_forced_for_diarization") else "No"),
        ("Worker diarización", f"persistente · job {metricas.get('diarization_worker_job_index') or '—'} · motor {'reutilizado' if metricas.get('diarization_engine_reused') else 'inicializado'}"),
        ("Preparación modelos diar.", _fmt_s(metricas.get("diarization_model_prepare_seconds"))),
        ("Inicialización motor diar.", _fmt_s(metricas.get("diarization_engine_init_seconds"))),
        ("Decodificación diar.", _fmt_s(metricas.get("diarization_decode_seconds"))),
        ("Proceso sherpa (wall)", _fmt_s(metricas.get("diarization_process_wall_seconds"))),
    ]
    if metricas.get("sherpa_internal_timing_available"):
        rows.extend([
            ("Sherpa segmentación", _fmt_s(internal.get("segmentation_seconds"))),
            ("Sherpa embeddings", _fmt_s(internal.get("embedding_seconds"))),
            ("Sherpa clustering", _fmt_s(internal.get("clustering_seconds"))),
            ("Sherpa total interno", _fmt_s(internal.get("sherpa_total_seconds"))),
        ])
    else:
        rows.append(("Timers internos sherpa", "no capturados por el backend/SO; se conserva medición wall"))
    if valid:
        rows.append(("Validación conteo hablantes", f"{valid.get('status', '—')} · estimación acústica {valid.get('estimated_speakers', '—')} · ground truth {'sí' if valid.get('ground_truth_available') else 'no'}"))

    for k, v in rows:
        cells = tabla.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v)
    doc.save(str(ruta))
