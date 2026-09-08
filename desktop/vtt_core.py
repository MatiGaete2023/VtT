#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Nucleo reutilizable de las mejoras VtT.

Este modulo no depende de Tkinter ni de Whisper. Mantiene separados los datos
tecnicos (segmentos finos) de las salidas de lectura (bloques agrupados), para
que SRT/VTT conserven precision mientras TXT/MD/DOCX sean legibles.
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple


PERFILES: Dict[str, Dict[str, Any]] = {
    "Rapido": {
        "batched": True,
        "batch_size": 8,
        "beam_size": 1,
        "descripcion": "Prioriza velocidad; menor busqueda durante la decodificacion.",
    },
    "Equilibrado": {
        "batched": True,
        "batch_size": 8,
        "beam_size": 5,
        "descripcion": "Recomendado: batching con la busqueda normal de Whisper.",
    },
    "Preciso": {
        "batched": False,
        "batch_size": 1,
        "beam_size": 5,
        "descripcion": "Ruta secuencial conservadora para audios dificiles.",
    },
}

EXTS_OFICIALES = (
    ".mp3", ".wav", ".m4a", ".ogg", ".flac", ".mp4", ".aac", ".wma",
    ".opus", ".webm", ".mkv", ".avi", ".mov", ".m4v", ".mpeg", ".mpg",
    ".3gp", ".ts", ".m2ts", ".aif", ".aiff",
)

FIN_ORACION_RE = re.compile(r"[.!?…][\"'”’)]*$")
ESPACIOS_RE = re.compile(r"\s+")


def _valor(obj: Any, nombre: str, defecto: Any = None) -> Any:
    if isinstance(obj, Mapping):
        return obj.get(nombre, defecto)
    return getattr(obj, nombre, defecto)


def limpiar_texto(texto: str) -> str:
    return ESPACIOS_RE.sub(" ", (texto or "").strip())


def segmento_a_dict(segmento: Any, indice: int = 0) -> Dict[str, Any]:
    """Convierte Segment de faster-whisper (o mapping de prueba) a JSON estable."""
    palabras = []
    for p in (_valor(segmento, "words", None) or []):
        palabras.append({
            "start": _valor(p, "start"),
            "end": _valor(p, "end"),
            "word": _valor(p, "word", _valor(p, "text", "")),
            "probability": _valor(p, "probability"),
        })
    return {
        "id": int(_valor(segmento, "id", indice) or indice),
        "start": float(_valor(segmento, "start", 0.0) or 0.0),
        "end": float(_valor(segmento, "end", 0.0) or 0.0),
        "text": limpiar_texto(str(_valor(segmento, "text", "") or "")),
        "words": palabras,
        "speaker": _valor(segmento, "speaker"),
        "avg_logprob": _valor(segmento, "avg_logprob"),
        "no_speech_prob": _valor(segmento, "no_speech_prob"),
        "compression_ratio": _valor(segmento, "compression_ratio"),
    }


def segmentos_a_dicts(segmentos: Iterable[Any]) -> List[Dict[str, Any]]:
    return [segmento_a_dict(s, i) for i, s in enumerate(segmentos)]


def _solapamiento(a0: float, a1: float, b0: float, b1: float) -> float:
    return max(0.0, min(a1, b1) - max(a0, b0))


def asignar_hablantes(
    segmentos: Sequence[Mapping[str, Any]],
    diarizacion: Sequence[Mapping[str, Any]],
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Asigna a cada segmento el hablante con mayor solapamiento temporal."""
    turnos = sorted(
        [
            {
                "start": float(t.get("start", 0.0)),
                "end": float(t.get("end", 0.0)),
                "speaker": int(t.get("speaker", 0)),
            }
            for t in diarizacion
            if float(t.get("end", 0.0)) > float(t.get("start", 0.0))
        ],
        key=lambda x: (x["start"], x["end"]),
    )
    ids = sorted({t["speaker"] for t in turnos})
    mapa_nombres = {sid: f"Persona {i + 1}" for i, sid in enumerate(ids)}

    salida: List[Dict[str, Any]] = []
    for original in segmentos:
        seg = dict(original)
        s0, s1 = float(seg.get("start", 0.0)), float(seg.get("end", 0.0))
        por_hablante: Dict[int, float] = {}
        for turno in turnos:
            ov = _solapamiento(s0, s1, turno["start"], turno["end"])
            if ov > 0:
                sid = turno["speaker"]
                por_hablante[sid] = por_hablante.get(sid, 0.0) + ov
        elegido: Optional[int] = None
        if por_hablante:
            elegido = max(por_hablante, key=por_hablante.get)
        else:
            medio = (s0 + s1) / 2.0
            candidatos = [t for t in turnos if t["start"] <= medio <= t["end"]]
            if candidatos:
                elegido = candidatos[0]["speaker"]
        seg["speaker_id"] = elegido
        seg["speaker"] = mapa_nombres.get(elegido) if elegido is not None else None
        salida.append(seg)

    speakers = [
        {"id": f"speaker_{sid:02d}", "numeric_id": sid, "display_name": mapa_nombres[sid]}
        for sid in ids
    ]
    return salida, speakers


def _confianza_bloque(segmentos: Sequence[Mapping[str, Any]]) -> Dict[str, Optional[float]]:
    def promedio(campo: str) -> Optional[float]:
        vals = [float(s[campo]) for s in segmentos if s.get(campo) is not None]
        return sum(vals) / len(vals) if vals else None

    return {
        "avg_logprob": promedio("avg_logprob"),
        "no_speech_prob": promedio("no_speech_prob"),
        "compression_ratio": promedio("compression_ratio"),
    }


def bloque_requiere_revision(bloque: Mapping[str, Any]) -> bool:
    conf = bloque.get("confidence") or {}
    logprob = conf.get("avg_logprob")
    no_speech = conf.get("no_speech_prob")
    compression = conf.get("compression_ratio")
    return bool(
        (logprob is not None and float(logprob) < -1.0)
        or (no_speech is not None and float(no_speech) > 0.6)
        or (compression is not None and float(compression) > 2.4)
    )


def agrupar_segmentos(
    segmentos: Sequence[Mapping[str, Any]],
    max_segundos: float = 35.0,
    max_caracteres: int = 650,
    pausa_corte: float = 1.4,
    min_caracteres_oracion: int = 180,
) -> List[Dict[str, Any]]:
    """Agrupa segmentos tecnicos en bloques de lectura."""
    limpios = [dict(s) for s in segmentos if limpiar_texto(str(s.get("text", "")))]
    if not limpios:
        return []
    bloques: List[Dict[str, Any]] = []
    actuales: List[Dict[str, Any]] = []

    def cerrar() -> None:
        if not actuales:
            return
        texto = limpiar_texto(" ".join(str(s.get("text", "")) for s in actuales))
        bloque = {
            "start": float(actuales[0].get("start", 0.0)),
            "end": float(actuales[-1].get("end", 0.0)),
            "speaker": actuales[0].get("speaker"),
            "text": texto,
            "segment_ids": [s.get("id") for s in actuales],
            "confidence": _confianza_bloque(actuales),
        }
        bloque["review_required"] = bloque_requiere_revision(bloque)
        bloques.append(bloque)
        actuales.clear()

    for seg in limpios:
        if not actuales:
            actuales.append(seg)
            continue
        previo = actuales[-1]
        speaker_cambia = seg.get("speaker") != previo.get("speaker") and (
            seg.get("speaker") is not None or previo.get("speaker") is not None
        )
        pausa = max(0.0, float(seg.get("start", 0.0)) - float(previo.get("end", 0.0)))
        duracion_si_agrega = float(seg.get("end", 0.0)) - float(actuales[0].get("start", 0.0))
        texto_actual = limpiar_texto(" ".join(str(x.get("text", "")) for x in actuales))
        texto_nuevo = limpiar_texto(str(seg.get("text", "")))
        chars_si_agrega = len(texto_actual) + 1 + len(texto_nuevo)
        termina_oracion = bool(FIN_ORACION_RE.search(str(previo.get("text", "")).strip()))
        corte_oracion = termina_oracion and len(texto_actual) >= min_caracteres_oracion
        if (speaker_cambia or pausa >= pausa_corte or duracion_si_agrega > max_segundos
                or chars_si_agrega > max_caracteres or corte_oracion):
            cerrar()
        actuales.append(seg)
    cerrar()
    return bloques


def ts_simple(t: float) -> str:
    total = max(0, int(float(t)))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def ts_srt(t: float) -> str:
    ms = max(0, int(round(float(t) * 1000)))
    h, ms = divmod(ms, 3600000)
    m, ms = divmod(ms, 60000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def ts_vtt(t: float) -> str:
    return ts_srt(t).replace(",", ".")


def texto_bloques(bloques: Sequence[Mapping[str, Any]], incluir_tiempos: bool = True) -> str:
    salida: List[str] = []
    for b in bloques:
        cabecera = []
        if incluir_tiempos:
            cabecera.append(f"[{ts_simple(float(b.get('start', 0.0)))}]")
        if b.get("speaker"):
            cabecera.append(str(b["speaker"]))
        if b.get("review_required"):
            cabecera.append("[REVISAR]")
        if cabecera:
            salida.append(" ".join(cabecera))
        salida.append(limpiar_texto(str(b.get("text", ""))))
        salida.append("")
    return "\n".join(salida).rstrip() + "\n"


def markdown_bloques(archivo: str, modelo: str, idioma: str,
                      bloques: Sequence[Mapping[str, Any]], metricas: Mapping[str, Any]) -> str:
    lineas = [
        f"# Transcripción: {Path(archivo).name}", "", f"- **Modelo:** {modelo}",
        f"- **Idioma:** {idioma or 'auto'}",
        f"- **Duración:** {ts_simple(float(metricas.get('audio_seconds', 0.0) or 0.0))}",
        f"- **Procesamiento:** {ts_simple(float(metricas.get('processing_seconds', 0.0) or 0.0))}",
        f"- **Velocidad:** {float(metricas.get('speed_x', 0.0) or 0.0):.2f}× tiempo real",
        "", "## Transcripción", "",
    ]
    for b in bloques:
        speaker = f" · **{b['speaker']}**" if b.get("speaker") else ""
        revisar = " · ⚠ **REVISAR**" if b.get("review_required") else ""
        lineas.append(f"### [{ts_simple(float(b.get('start', 0.0)))}]{speaker}{revisar}")
        lineas.append("")
        lineas.append(limpiar_texto(str(b.get("text", ""))))
        lineas.append("")
    return "\n".join(lineas).rstrip() + "\n"


def srt_segmentos(segmentos: Sequence[Mapping[str, Any]]) -> str:
    out: List[str] = []
    for i, s in enumerate(segmentos, 1):
        out.append(str(i))
        out.append(f"{ts_srt(float(s['start']))} --> {ts_srt(float(s['end']))}")
        prefijo = f"{s.get('speaker')}: " if s.get("speaker") else ""
        out.append(prefijo + limpiar_texto(str(s.get("text", ""))))
        out.append("")
    return "\n".join(out)


def vtt_segmentos(segmentos: Sequence[Mapping[str, Any]]) -> str:
    out = ["WEBVTT", ""]
    for s in segmentos:
        out.append(f"{ts_vtt(float(s['start']))} --> {ts_vtt(float(s['end']))}")
        prefijo = f"{s.get('speaker')}: " if s.get("speaker") else ""
        out.append(prefijo + limpiar_texto(str(s.get("text", ""))))
        out.append("")
    return "\n".join(out)


def calcular_metricas(audio_seconds: float, processing_seconds: float, **extras: Any) -> Dict[str, Any]:
    audio = max(0.0, float(audio_seconds or 0.0))
    proc = max(0.0, float(processing_seconds or 0.0))
    rtf = (proc / audio) if audio > 0 else None
    speed_x = (audio / proc) if proc > 0 else None
    d = {"audio_seconds": audio, "processing_seconds": proc, "rtf": rtf, "speed_x": speed_x}
    d.update(extras)
    return d


def documento_json(archivo: str, modelo: str, idioma: str, texto: str,
                    segmentos: Sequence[Mapping[str, Any]], bloques: Sequence[Mapping[str, Any]],
                    speakers: Sequence[Mapping[str, Any]], metricas: Mapping[str, Any],
                    perfil: str, glosario: str) -> Dict[str, Any]:
    return {
        "schema_version": 2,
        "source": Path(archivo).name,
        "model": modelo,
        "language": idioma or "auto",
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "profile": perfil,
        "hotwords": [x.strip() for x in re.split(r"[,;\n]+", glosario or "") if x.strip()],
        "text": texto,
        "metrics": dict(metricas),
        "segments": [dict(s) for s in segmentos],
        "reading_blocks": [dict(b) for b in bloques],
        "speakers": [dict(s) for s in speakers],
        "review": {
            "status": "machine_generated", "edited": False,
            "flagged_blocks": sum(1 for b in bloques if b.get("review_required")),
        },
    }


def json_texto(documento: Mapping[str, Any]) -> str:
    return json.dumps(documento, ensure_ascii=False, indent=2) + "\n"


def escribir_docx(ruta: Path, archivo: str, modelo: str, idioma: str,
                   bloques: Sequence[Mapping[str, Any]], metricas: Mapping[str, Any]) -> None:
    """Escribe DOCX literal y estructurado. No reescribe el contenido ASR."""
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("TRANSCRIPCIÓN", level=0)
    tabla = doc.add_table(rows=0, cols=2)
    datos = [
        ("Archivo", Path(archivo).name), ("Modelo", modelo), ("Idioma", idioma or "auto"),
        ("Duración", ts_simple(float(metricas.get("audio_seconds", 0.0) or 0.0))),
        ("Procesamiento", ts_simple(float(metricas.get("processing_seconds", 0.0) or 0.0))),
        ("Velocidad", f"{float(metricas.get('speed_x', 0.0) or 0.0):.2f}× tiempo real"),
    ]
    for k, v in datos:
        c = tabla.add_row().cells
        c[0].text = k
        c[1].text = str(v)
    doc.add_paragraph()
    for b in bloques:
        p = doc.add_paragraph()
        run = p.add_run(f"[{ts_simple(float(b.get('start', 0.0)))}]")
        run.bold = True
        if b.get("speaker"):
            r = p.add_run(f"  {b['speaker']}")
            r.bold = True
        if b.get("review_required"):
            r = p.add_run("  [REVISAR]")
            r.bold = True
        doc.add_paragraph(limpiar_texto(str(b.get("text", ""))))
    doc.save(str(ruta))


def glosario_a_hotwords(glosario: str) -> Optional[str]:
    partes = [x.strip() for x in re.split(r"[,;\n]+", glosario or "") if x.strip()]
    return ", ".join(dict.fromkeys(partes)) or None


def tiene_extension_conocida(ruta: str) -> bool:
    return Path(ruta).suffix.lower() in EXTS_OFICIALES


def probar_pista_audio(ruta: str) -> Tuple[bool, str]:
    """Valida con PyAV que exista al menos una pista de audio decodificable."""
    try:
        import av
        with av.open(ruta) as cont:
            pistas = [s for s in cont.streams if s.type == "audio"]
            if not pistas:
                return False, "El archivo no contiene una pista de audio."
            return True, f"Audio detectado: {pistas[0].codec_context.name or 'codec desconocido'}"
    except Exception as exc:
        return False, f"No se pudo abrir el audio/video: {exc}"


def extraer_audio_revision(ruta: str, inicio: float, segundos: float = 12.0, rate: int = 16000):
    """Devuelve numpy float32 mono para reproducir una ventana desde un timestamp."""
    import av
    import numpy as np
    inicio = max(0.0, float(inicio))
    segundos = max(0.5, float(segundos))
    piezas = []
    con = av.open(ruta)
    try:
        stream = next((s for s in con.streams if s.type == "audio"), None)
        if stream is None:
            raise ValueError("El archivo no contiene audio")
        if stream.time_base:
            target = int(inicio / float(stream.time_base))
            con.seek(target, stream=stream, backward=True, any_frame=False)
        resampler = av.AudioResampler(format="fltp", layout="mono", rate=rate)
        total_obj = int(segundos * rate)
        recogidas = 0
        for frame in con.decode(stream):
            tiempo = float(frame.pts * frame.time_base) if frame.pts is not None and frame.time_base else None
            if tiempo is not None and tiempo + float(frame.duration or 0) * float(frame.time_base or 0) < inicio:
                continue
            for rf in resampler.resample(frame):
                arr = rf.to_ndarray()
                mono = arr[0] if arr.ndim > 1 else arr
                mono = np.asarray(mono, dtype=np.float32)
                piezas.append(mono)
                recogidas += len(mono)
                if recogidas >= total_obj:
                    break
            if recogidas >= total_obj:
                break
        if not piezas:
            raise ValueError("No se pudo decodificar audio desde ese punto")
        data = np.concatenate(piezas)[:total_obj]
        return np.ascontiguousarray(data), rate
    finally:
        con.close()
