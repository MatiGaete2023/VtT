#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reportes V5.1: confianza por identidad, wall por pasada y advertencias."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Mapping, Sequence

from docx import Document

import vtt_reporting_v5 as v5
import vtt_validation_v51 as validation51


def documento_json_detallado(
    archivo: str, modelo: str, idioma: str, texto: str,
    segmentos: Sequence[Mapping[str, Any]], bloques: Sequence[Mapping[str, Any]],
    speakers: Sequence[Mapping[str, Any]], metricas: Mapping[str, Any],
    perfil: str, glosario: str,
) -> Dict[str, Any]:
    doc = v5.documento_json_detallado(
        archivo, modelo, idioma, texto, segmentos, bloques,
        speakers, metricas, perfil, glosario,
    )
    doc["schema_version"] = max(6, int(doc.get("schema_version", 0) or 0))
    diar = doc.setdefault("diarization", {})
    diar["identity_verification"] = metricas.get("identity_verification") or {}
    diar["identity_refinement"] = metricas.get("identity_refinement") or {}
    diar["local_change_scan"] = metricas.get("identity_local_scan") or {}
    diar["identity_consistency"] = metricas.get("identity_consistency") or {}
    return doc


def _fmt_s(value: Any) -> str:
    try:
        return f"{float(value or 0.0):.2f} s"
    except (TypeError, ValueError):
        return "—"


def _find_row(table, label: str):
    for row in table.rows:
        if row.cells and row.cells[0].text.strip() == label:
            return row
    return None


def escribir_docx_detallado(
    ruta: Path, archivo: str, modelo: str, idioma: str,
    bloques: Sequence[Mapping[str, Any]], metricas: Mapping[str, Any], perfil: str,
) -> None:
    v5.escribir_docx_detallado(ruta, archivo, modelo, idioma, bloques, metricas, perfil)
    doc = Document(str(ruta))
    if not doc.tables:
        doc.save(str(ruta)); return
    tabla = doc.tables[0]
    valid = metricas.get("speaker_count_validation") or {}
    row = _find_row(tabla, "Hablantes")
    if row is not None and metricas.get("speaker_mode") == "Auto":
        n = int(metricas.get("speaker_detected", 0) or 0)
        row.cells[1].text = f"Auto · estimación: {n} · {validation51.etiqueta_confianza(valid)}"

    rows = []
    for p in metricas.get("sherpa_pass_timings") or []:
        rows.append((f"Sherpa pasada {int(p.get('pass', 0) or 0)} (wall)", _fmt_s(p.get("process_wall_seconds"))))

    iv = metricas.get("identity_verification") or {}
    consistency = metricas.get("identity_consistency") or {}
    refinement = metricas.get("identity_refinement") or {}
    scan = metricas.get("identity_local_scan") or {}
    if iv.get("enabled"):
        rows.extend([
            ("Control identidad V5.1", "Sí · prototipos acústicos conservadores"),
            ("Verificación identidad", _fmt_s(metricas.get("identity_wall_seconds"))),
            ("Embeddings identidad", f"{int(metricas.get('identity_embedding_calls', 0) or 0)} cálculo(s) · {_fmt_s(metricas.get('identity_embedding_compute_seconds'))}"),
            ("Refinamiento identidad", f"{int(refinement.get('reassigned_turns', 0) or 0)} turno(s) reasignado(s) · {int(refinement.get('new_identities', 0) or 0)} identidad(es) nueva(s)"),
            ("Escaneo local", f"{int(scan.get('scanned_turns', 0) or 0)} turno(s) largo(s) · {int(scan.get('applied_changes', 0) or 0)} cambio(s) aplicado(s)"),
            ("Consistencia identidad", str(consistency.get("overall_confidence", "sin_datos"))),
        ])
        for sp in consistency.get("speakers", []) or []:
            name = sp.get("display_name") or f"ID raw {sp.get('speaker')}"
            detail = f"{sp.get('confidence', 'sin_datos')} · {int(sp.get('turns', 0) or 0)} turno(s) · {float(sp.get('total_seconds', 0.0) or 0.0):.1f} s"
            if sp.get("prototype_similarity_median") is not None:
                detail += f" · similitud prototipo mediana {float(sp['prototype_similarity_median']):.2f}"
            rows.append((f"Identidad {name}", detail))

    if str(metricas.get("device", "")) == "cpu" and metricas.get("diarization_profile") == "Precisa":
        rows.append(("Rendimiento diarización", "Precisa prioriza resolución temporal y puede ser muy lenta en CPU; Equilibrada es el perfil recomendado para uso habitual."))

    if valid:
        rowv = _find_row(tabla, "Validación conteo hablantes")
        if rowv is not None:
            rowv.cells[1].text = (
                f"{valid.get('status', '—')} · estimación acústica {valid.get('estimated_speakers', '—')} · "
                f"ground truth {'sí' if valid.get('ground_truth_available') else 'no'} · identidad {valid.get('identity_confidence', 'sin_datos')}"
            )
    for k, v in rows:
        cells = tabla.add_row().cells; cells[0].text = str(k); cells[1].text = str(v)
    doc.save(str(ruta))
